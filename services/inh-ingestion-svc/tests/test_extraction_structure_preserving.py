"""Structure-preserving DOCX/PDF extraction tests (#389).

Before #389, `_extract_docx_text` returned only
`[p.text for p in doc.paragraphs]` (top-level paragraphs, no tables, no
heading structure, no rendered list numbers) and `_extract_pdf_text` did a
flat `page.extract_text()` join with no post-processing. These tests build
fixtures programmatically (python-docx for DOCX; a small hand-written PDF
content stream, since no PDF-generation library is available in this
service's dependency set -- see `_build_pdf` below) and pin the new,
structure-preserving contract described in `extract.py`'s
`_extract_docx_text`/`_extract_pdf_text` docstrings.
"""

import io

import docx
import pytest
from docx.oxml.ns import qn

from src.temporal.activities.extract import _extract_docx_text, _extract_pdf_text


@pytest.fixture(autouse=True)
def cleanup_test_data():
    """Override the DB-backed root autouse fixture (see
    `test_extraction_by_type.py`'s identical override) -- offline, no
    PostgreSQL required."""
    yield


# ---------------------------------------------------------------------------
# DOCX fixture helpers -- numbering.xml has no public python-docx write API
# (see `docx_numbering.py`'s module docstring and `test_docx_numbering.py`'s
# identical helpers), so these build it directly via the OOXML element tree.
# ---------------------------------------------------------------------------


def _set_val(parent, tag: str, value: str) -> None:
    child = parent.makeelement(qn(tag), {})
    child.set(qn("w:val"), value)
    parent.append(child)


def _add_abstract_num(numbering_element, abstract_id: str, levels: list[tuple[str, str, str]]):
    abstract_num = numbering_element.makeelement(qn("w:abstractNum"), {})
    abstract_num.set(qn("w:abstractNumId"), abstract_id)
    for ilvl, (num_fmt, lvl_text, start) in enumerate(levels):
        lvl = abstract_num.makeelement(qn("w:lvl"), {})
        lvl.set(qn("w:ilvl"), str(ilvl))
        _set_val(lvl, "w:start", start)
        _set_val(lvl, "w:numFmt", num_fmt)
        _set_val(lvl, "w:lvlText", lvl_text)
        abstract_num.append(lvl)
    numbering_element.append(abstract_num)


def _add_num(numbering_element, num_id: str, abstract_id: str):
    num = numbering_element.makeelement(qn("w:num"), {})
    num.set(qn("w:numId"), num_id)
    _set_val(num, "w:abstractNumId", abstract_id)
    numbering_element.append(num)


def _set_direct_num_pr(paragraph, num_id: str, ilvl: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    _set_val(num_pr, "w:ilvl", str(ilvl))
    _set_val(num_pr, "w:numId", str(num_id))


def _set_style_num_pr(style, num_id: str, ilvl: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    _set_val(num_pr, "w:ilvl", str(ilvl))
    _set_val(num_pr, "w:numId", str(num_id))


def _save(document: docx.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDocxHeadings:
    def test_heading_styles_render_as_their_own_plain_lines(self):
        """Headings (Title and Heading 1..6) keep their own line, without a
        markdown "#" marker: DOCX is chunked as prose (which ignores "#"),
        numbered_sections detects sections by numbering, and the marker
        measurably lowered retrieval (retrieval-eval gate, recall@5 0.88 ->
        0.84)."""
        document = docx.Document()
        document.add_heading("Product Handbook", level=0)  # "Title" style
        document.add_heading("Definitions", level=1)
        document.add_heading("Payment Terms", level=2)
        document.add_paragraph("Plain body text, no heading style.")

        text = _extract_docx_text(_save(document), "sample.docx")
        lines = text.splitlines()

        for heading in ("Product Handbook", "Definitions", "Payment Terms"):
            assert heading in lines, f"{heading!r} must be a line of its own"
        assert not any(line.startswith("#") for line in lines)
        assert "Plain body text, no heading style." in text

    def test_all_six_heading_levels_stay_plain_lines(self):
        document = docx.Document()
        for level in range(1, 7):
            document.add_heading(f"Heading level {level}", level=level)

        lines = _extract_docx_text(_save(document), "sample.docx").splitlines()

        for level in range(1, 7):
            assert f"Heading level {level}" in lines


class TestDocxNumbering:
    def test_multi_level_numbering_with_counter_resets(self):
        """1 / 1.1 / (a) / (i), then back up to 2 -- the #389 acceptance
        example, including that returning to a shallower level resets the
        deeper levels rather than continuing them."""
        document = docx.Document()
        numbering_element = document.part.numbering_part.element
        _add_abstract_num(
            numbering_element,
            "60",
            [
                ("decimal", "%1", "1"),
                ("decimal", "%1.%2", "1"),
                ("lowerLetter", "(%3)", "1"),
                ("lowerRoman", "(%4)", "1"),
            ],
        )
        _add_num(numbering_element, "60", "60")

        texts_and_levels = [
            ("Definitions", 0),
            ("The team ships releases within 30 days", 1),
            ("any breach of this clause shall be notified", 2),
            ("in writing to the owner", 3),
            ("Payment Terms", 0),
            ("Invoices are due net 30", 1),
        ]
        for body_text, ilvl in texts_and_levels:
            paragraph = document.add_paragraph(body_text)
            _set_direct_num_pr(paragraph, "60", ilvl)

        text = _extract_docx_text(_save(document), "sample.docx")

        assert "1 Definitions" in text
        assert "1.1 The team ships releases within 30 days" in text
        assert "(a) any breach of this clause shall be notified" in text
        assert "(i) in writing to the owner" in text
        # Counter reset: the second top-level item is "2", not "3", and its
        # own sub-level restarts at ".1", not continuing "1.2"/".2".
        assert "2 Payment Terms" in text
        assert "2.1 Invoices are due net 30" in text

    def test_bullets_render_as_dash(self):
        document = docx.Document()
        numbering_element = document.part.numbering_part.element
        _add_abstract_num(numbering_element, "61", [("bullet", "", "1")])
        _add_num(numbering_element, "61", "61")

        for item in ("First bullet item", "Second bullet item"):
            paragraph = document.add_paragraph(item)
            _set_direct_num_pr(paragraph, "61", 0)

        text = _extract_docx_text(_save(document), "sample.docx")

        assert "- First bullet item" in text
        assert "- Second bullet item" in text

    def test_style_inherited_num_pr(self):
        """A paragraph using a list-numbered STYLE, with no numPr of its own
        -- the template-driven case (#389 requirement)."""
        document = docx.Document()
        numbering_element = document.part.numbering_part.element
        _add_abstract_num(numbering_element, "62", [("decimal", "%1.", "1")])
        _add_num(numbering_element, "62", "62")

        list_style = document.styles.add_style(
            "NumberedParagraph", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH
        )
        _set_style_num_pr(list_style, "62", 0)

        document.add_paragraph("First styled item", style="NumberedParagraph")
        document.add_paragraph("Second styled item", style="NumberedParagraph")

        text = _extract_docx_text(_save(document), "sample.docx")

        assert "1. First styled item" in text
        assert "2. Second styled item" in text


class TestDocxTablesInDocumentOrder:
    def test_table_renders_as_markdown_and_stays_in_document_order(self):
        """A table between two paragraphs must appear BETWEEN them in the
        extracted text (#389: document-order walk, not paragraphs-then-
        tables) and render as a GitHub markdown table with pipes escaped."""
        document = docx.Document()
        document.add_paragraph("Before the table.")
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Name"
        table.rows[0].cells[1].text = "Value"
        table.rows[1].cells[0].text = "A|B"
        table.rows[1].cells[1].text = "42"
        document.add_paragraph("After the table.")

        text = _extract_docx_text(_save(document), "sample.docx")

        before_pos = text.index("Before the table.")
        table_pos = text.index("| Name | Value |")
        after_pos = text.index("After the table.")
        assert before_pos < table_pos < after_pos

        assert "| Name | Value |" in text
        assert "| --- | --- |" in text
        # The literal pipe inside a cell's own text must be escaped so it
        # can't be misread as an extra column delimiter.
        assert "A\\|B" in text
        assert "| 42 |" in text


class TestDocxAcceptanceSyntheticDocument:
    """#389 acceptance criterion: a ~20-section synthetic document must
    survive extraction with every section boundary intact (target >= 19/20).
    """

    def _build_synthetic_document(self, section_count: int = 20) -> bytes:
        document = docx.Document()
        numbering_element = document.part.numbering_part.element
        _add_abstract_num(numbering_element, "70", [("decimal", "%1.", "1")])
        _add_num(numbering_element, "70", "70")

        document.add_heading("Synthetic Handbook", level=0)
        for n in range(1, section_count + 1):
            heading = document.add_paragraph(f"Section {n} Heading")
            _set_direct_num_pr(heading, "70", 0)
            document.add_paragraph(
                f"This is the body text for section {n}, describing the "
                f"obligations relevant to clause {n}."
            )
        return _save(document)

    def test_every_section_boundary_survives_extraction(self):
        content = self._build_synthetic_document(section_count=20)
        text = _extract_docx_text(content, "synthetic-document.docx")

        intact = 0
        for n in range(1, 21):
            marker = f"{n}. Section {n} Heading"
            body = f"body text for section {n},"
            if marker in text and body in text:
                intact += 1

        assert intact >= 19, f"only {intact}/20 section boundaries survived intact"


# ---------------------------------------------------------------------------
# PDF fixtures -- hand-written content stream (no reportlab/fpdf in this
# service's dependency set), following the same minimal single-page PDF
# shape as the bundled `docs/examples/sample-documents/sample.pdf` fixture.
# ---------------------------------------------------------------------------


def _pdf_escape(line: str) -> str:
    return line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _build_pdf(lines: list[str], leading: int = 14) -> bytes:
    """A minimal one-page PDF whose content stream renders `lines` one per
    text line via `Tj`/`T*`, mirroring exactly how `sample.pdf` (see
    `docs/examples/README.md`'s generation recipe) was built by hand."""
    content_lines = ["BT", "/F1 10 Tf", "50 740 Td", f"{leading} TL"]
    for index, line in enumerate(lines):
        escaped = _pdf_escape(line)
        if index > 0:
            content_lines.append("T*")
        content_lines.append(f"({escaped}) Tj")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_offset).encode()
        + b"\n%%EOF"
    )
    return bytes(out)


class TestPdfLineStructure:
    def test_numbered_heading_lines_stay_on_their_own_line(self):
        lines = [
            "1. Definitions",
            "1.1 The team ships releases within 30 days.",
            "(a) any breach of this clause shall be notified",
            "(i) in writing to the owner",
            "2. Payment Terms",
        ]
        text = _extract_pdf_text(_build_pdf(lines))
        extracted_lines = text.split("\n")
        for expected in lines:
            assert expected in extracted_lines, (
                f"expected line {expected!r} to survive as its own line in " f"{extracted_lines!r}"
            )

    def test_tight_line_leading_still_keeps_lines_separate(self):
        """A dense document layout (small leading) must not merge adjacent
        numbered lines into one (#389: "line breaks are preserved")."""
        lines = ["3.1 First tightly-spaced clause.", "3.2 Second tightly-spaced clause."]
        text = _extract_pdf_text(_build_pdf(lines, leading=10))
        extracted_lines = text.split("\n")
        assert "3.1 First tightly-spaced clause." in extracted_lines
        assert "3.2 Second tightly-spaced clause." in extracted_lines


class TestPdfHyphenationAndWhitespaceNormalization:
    def test_line_wrap_hyphen_is_rejoined(self):
        # pypdf's own line-break detection already turns this into two
        # `extract_text()` lines joined by "\n" -- exactly the wrap-artifact
        # shape `_normalize_pdf_page_text` targets.
        text = _extract_pdf_text(_build_pdf(["This is informa-", "tion about the release."]))
        assert "information about the release." in text
        assert "informa-\ntion" not in text

    def test_numbered_marker_is_never_dehyphenated(self):
        """A numbered marker never contains a lowercase-letter/hyphen shape
        the dehyphenation regex could match, but this pins that explicitly:
        section numbers must survive byte-for-byte."""
        text = _extract_pdf_text(
            _build_pdf(["1.1 Confidentiality obligations survive termination."])
        )
        assert "1.1 Confidentiality obligations survive termination." in text
