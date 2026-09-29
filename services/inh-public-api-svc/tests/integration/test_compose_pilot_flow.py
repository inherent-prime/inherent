"""Live pilot-flow E2E: a vertical pack, one bound workspace, end to end.

Everything a hand-onboarded pilot depends on, proven against a RUNNING stack
with real Postgres / Mongo / Weaviate / Temporal / embeddings, one ordered
step at a time (a. ingest .. g. purge):

    a. a synthetic .docx (numbered sections, 1 / 1.1 / (a) outline numbering, a
       table) is ingested with a ``source_url`` and reaches ``processed``;
    b. its chunks split at the numbered sections, carry the parent heading and
       the pack's rule-derived tags;
    c. MCP Streamable HTTP (``/mcp``, API-key auth): ``tools/list`` shows the
       pack's profile tool; calling it with a tag filter returns text, tags, the
       section heading and the ``source_url``;
    d. that call left an audit event (principal / surface / tool name / the
       returned chunk ids) in Mongo ``audit_logs``;
    e. a second document reusing one section bumps the OLDER chunk's
       ``reuse_count``;
    f. workspace membership: a ``viewer`` can search the owner's documents
       (owner-as-tenant) but not upload; once removed, search is denied;
    g. ``POST /admin/workspaces/{id}/purge`` -> poll -> the receipt reports zero
       residue in every store.

The test is pack-agnostic: it reads the pack's manifest from the directory in
``E2E_VERTICAL_PACK_DIR`` and derives the tool name, filter field and the words
that trigger each tag value from the pack's own files. No domain vocabulary
lives here. Without ``E2E_VERTICAL_PACK_DIR`` it skips.

The stack must have been booted with that pack mounted and this workspace bound
to it (see ``docker-compose.pilot-e2e.yml``, which does exactly that)::

    E2E_VERTICAL_PACK_DIR=$PWD/services/inh-public-api-svc/tests/fixtures/handbook_pack \\
      docker compose -f docker-compose.yml -f docker-compose.pilot-e2e.yml up -d --build --wait
    make bootstrap
    cd services/inh-public-api-svc
    E2E_VERTICAL_PACK_DIR=... uv run pytest tests/integration/test_compose_pilot_flow.py -m compose --no-cov

Configuration (all have local defaults; override via env):
    E2E_VERTICAL_PACK_DIR      host path of the pack root (REQUIRED, else skip)
    E2E_PILOT_WORKSPACE_ID     default ws_pilot_e2e_001 (must match the stack's
                               WORKSPACE_VERTICAL_PACKS binding)
    PUBLIC_API_URL             default http://localhost:18000
    INGESTION_API_URL          default http://localhost:18002
    INGESTION_API_KEY          default dev-ingestion-key
    INTEGRATION_DATABASE_URL   default postgresql://postgres:postgres@localhost:15432/knowledge_base
    MONGODB_URI / MONGODB_DB_NAME  default mongodb://localhost:27018 / main
    E2E_S3_ENDPOINT            default http://localhost:19000 (AWS_ACCESS_KEY_ID /
                               AWS_SECRET_ACCESS_KEY / AWS_S3_BUCKET as the stack)
    INTEGRATION_TIMEOUT        seconds to wait for ingestion (default 180)
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import os
import re
import time
import uuid
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import boto3
import httpx
import pytest
from botocore.config import Config
from docx import Document
from inh_contracts.vertical_pack import Vertical, load_vertical
from pymongo import MongoClient

from tests.integration.test_compose_mcp import _structured_payload, mcp_http_session

pytestmark = [pytest.mark.compose, pytest.mark.integration, pytest.mark.slow]

PACK_DIR_ENV = os.environ.get("E2E_VERTICAL_PACK_DIR")

API_URL = os.environ.get("PUBLIC_API_URL", "http://localhost:18000").rstrip("/")
INGESTION_URL = os.environ.get("INGESTION_API_URL", "http://localhost:18002").rstrip("/")
INGESTION_KEY = os.environ.get("INGESTION_API_KEY", "dev-ingestion-key")
# Deliberately NOT plain DATABASE_URL: this test WRITES api_keys rows, and a
# developer shell often has DATABASE_URL pointed at some other (unit-test) database.
DATABASE_URL = os.environ.get(
    "INTEGRATION_DATABASE_URL", "postgresql://postgres:postgres@localhost:15432/knowledge_base"
)
MONGODB_URI = os.environ.get("MONGODB_URI", "mongodb://localhost:27018")
MONGODB_DB = os.environ.get("MONGODB_DB_NAME", "main")
S3_ENDPOINT = os.environ.get("E2E_S3_ENDPOINT", "http://localhost:19000")
S3_BUCKET = os.environ.get("AWS_S3_BUCKET", "inherent-documents")
TIMEOUT = int(os.environ.get("INTEGRATION_TIMEOUT", "180"))

WORKSPACE_ID = os.environ.get("E2E_PILOT_WORKSPACE_ID", "ws_pilot_e2e_001")

# Two principals in the pilot workspace. Fixed ids (not random) because the
# stack binds the WORKSPACE, not the users, and re-runs must upsert, not pile up.
OWNER_USER = "pilot-e2e-owner"
OWNER_KEY_ID = "pilot-e2e-owner-key"
OWNER_KEY = "ink_pilot_e2e_owner_key_001"
VIEWER_USER = "pilot-e2e-viewer"
VIEWER_KEY_ID = "pilot-e2e-viewer-key"
VIEWER_KEY = "ink_pilot_e2e_viewer_key_001"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Unique per run so a stale document from an aborted earlier run can never make
# an assertion pass or fail for the wrong reason.
RUN = uuid.uuid4().hex[:8]
SOURCE_URL_1 = f"https://files.example.test/pilot/{RUN}/first"
SOURCE_URL_2 = f"https://files.example.test/pilot/{RUN}/second"


# ---------------------------------------------------------------------------
# Deriving test content from the pack (no domain terms hard-coded here)
# ---------------------------------------------------------------------------


def _sample_for_pattern(pattern: str) -> str | None:
    """A short string that ``pattern`` (a tag rule) matches, or None.

    Rules are case-insensitive keywords/regexes. This expands the simple shapes
    packs use (``\\b`` anchors, ``(a|b)`` groups, ``x?`` optionals) into a
    literal sample and VERIFIES it against the real regex, so a shape it cannot
    expand is skipped rather than guessed.
    """
    sample = pattern.replace("\\b", "")
    sample = re.sub(r"\([^()]*\)\?", "", sample)  # optional group -> drop
    sample = re.sub(r"\(([^()|]*)(?:\|[^()]*)*\)", r"\1", sample)  # (a|b) -> a
    sample = re.sub(r".\?", "", sample)  # x? -> (nothing)
    sample = sample.replace("\\", "")
    try:
        return sample if re.search(pattern, sample, re.IGNORECASE) else None
    except re.error:
        return None


@dataclass
class PackPlan:
    """What the test needs to know about the pack, all read from its files."""

    pack: Vertical
    tool_name: str
    filter_field: str
    value_a: str  # tag value planted in two sections (used as the filter)
    value_b: str  # a different tag value planted in one section
    trigger_a: str
    trigger_b: str


def _plan_from_pack(root: Path) -> PackPlan:
    pack = load_vertical(root)
    tool = pack.tools[0]
    for field_name in tool.filters:
        spec = pack.tags.fields[field_name]
        if spec.type != "enum" or not spec.rules:
            continue
        triggers: dict[str, str] = {}
        for value, patterns in spec.rules.items():
            for pattern in patterns:
                sample = _sample_for_pattern(pattern)
                if sample:
                    triggers[value] = sample
                    break
        if len(triggers) >= 2:
            (value_a, trigger_a), (value_b, trigger_b) = list(triggers.items())[:2]
            plan = PackPlan(pack, tool.name, field_name, value_a, value_b, trigger_a, trigger_b)
            _assert_filler_is_neutral(pack)
            return plan
    raise AssertionError(
        f"pack {pack.manifest.name!r}: its first tool profile must expose at least one "
        "enum filter field with >= 2 rule-tagged values, so the flow can prove filtering"
    )


# Body text with no tag-rule trigger words for any pack. Verified against the
# loaded pack's rules so a pack whose rules match it fails loudly, not silently.
FILLER = "This paragraph describes the working arrangement between the two sides."
SUBITEM_1 = "first sub-item text that continues the numbered section above"
SUBITEM_2 = "second sub-item text that continues the numbered section above"


def _assert_filler_is_neutral(pack: Vertical) -> None:
    for field_name, spec in pack.tags.fields.items():
        for patterns in spec.rules.values():
            for pattern in patterns:
                for text in (
                    FILLER,
                    SUBITEM_1,
                    SUBITEM_2,
                    "General Provisions",
                    "Additional Terms",
                ):
                    assert not re.search(
                        pattern, text, re.IGNORECASE
                    ), f"filler {text!r} unexpectedly matches rule {pattern!r} of {field_name}"


def _sentence(trigger: str, marker: str) -> str:
    """A long-ish section body that contains exactly one rule trigger.

    Long on purpose: the section body, not its short heading, should dominate
    the chunk text the reuse step compares.
    """
    return (
        f"The sides confirm that {trigger} applies to every arrangement covered by "
        f"reference {marker}, and that the applicable terms are recorded in full in this "
        f"section so that a later reader can rely on the wording without further explanation."
    )


def _build_docx_first(plan: PackPlan) -> bytes:
    """Numbered sections, 1 / 1.1 / (a) outline numbering and a table."""
    doc = Document()
    doc.add_heading("1. General Provisions", level=1)
    doc.add_paragraph(FILLER)
    doc.add_paragraph(f"1.1 {_sentence(plan.trigger_a, f'A1-{RUN}')}")
    doc.add_paragraph(f"(a) {SUBITEM_1}")
    doc.add_paragraph(f"(b) {SUBITEM_2}")
    doc.add_paragraph(f"1.2 {_sentence(plan.trigger_b, f'B1-{RUN}')}")
    doc.add_heading("2. Additional Terms", level=1)
    doc.add_paragraph(f"2.1 {_sentence(plan.trigger_a, f'A2-{RUN}')}")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Detail"
    table.rows[1].cells[0].text = f"TBLCELL-{RUN}"
    table.rows[1].cells[1].text = "listed"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# The reused section below repeats the FIRST document's 1.1 section verbatim,
# under the same parent heading (a copied section arrives with its numbering and
# context; the same wording under a different heading embeds measurably
# further apart and falls below the engine's 0.92 cosine reuse threshold).
# Everything else in this document is deliberately unrelated.
def _build_docx_second(plan: PackPlan) -> bytes:
    doc = Document()
    doc.add_heading("1. General Provisions", level=1)
    doc.add_paragraph("Goods ship on the first business day of each calendar month.")
    doc.add_paragraph(f"1.1 {_sentence(plan.trigger_a, f'A1-{RUN}')}")  # <- reused section
    doc.add_paragraph("1.2 Late shipments are reported to the coordinator within two days.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Stack access helpers
# ---------------------------------------------------------------------------


def _pg_upsert_key(key_id: str, api_key: str, user_id: str) -> None:
    """Insert (or re-activate) a workspace-scoped API key, the way bootstrap.sh does."""
    import asyncpg

    async def _run() -> None:
        conn = await asyncpg.connect(DATABASE_URL)
        try:
            await conn.execute(
                """
                INSERT INTO api_keys
                  (key_id, key_hash, key_prefix, user_id, workspace_id, name, status,
                   permissions, rate_limit)
                VALUES ($1, $2, $3, $4, $5, $6, 'active',
                        '["read","write","search"]', 1000)
                ON CONFLICT (key_hash) DO UPDATE
                  SET status = 'active', user_id = EXCLUDED.user_id,
                      workspace_id = EXCLUDED.workspace_id, permissions = EXCLUDED.permissions
                """,
                key_id,
                hashlib.sha256(api_key.encode()).hexdigest(),
                api_key[:12],
                user_id,
                WORKSPACE_ID,
                f"pilot e2e {user_id}",
            )
        finally:
            await conn.close()

    asyncio.run(_run())


def _pg_delete_keys() -> None:
    import asyncpg

    async def _run() -> None:
        conn = await asyncpg.connect(DATABASE_URL)
        try:
            await conn.execute(
                "DELETE FROM api_keys WHERE key_hash = ANY($1::text[])",
                [hashlib.sha256(k.encode()).hexdigest() for k in (OWNER_KEY, VIEWER_KEY)],
            )
        finally:
            await conn.close()

    asyncio.run(_run())


def _pg_clear_purge_tombstone() -> None:
    """Forget that the pilot workspace id was ever purged.

    A purge is terminal by design: it leaves a ``workspace_purge_state`` row and
    the engine then refuses every new ingest for that workspace id. This test
    purges the same fixed id at the end of every run (step g), so a re-run must
    lift that marker first -- test-only bookkeeping on a throwaway workspace id.
    """
    import asyncpg

    async def _run() -> None:
        conn = await asyncpg.connect(DATABASE_URL)
        try:
            await conn.execute(
                "DELETE FROM workspace_purge_state WHERE workspace_id = $1", WORKSPACE_ID
            )
        finally:
            await conn.close()

    asyncio.run(_run())


def _headers(api_key: str) -> dict[str, str]:
    return {"X-API-Key": api_key, "X-Workspace-Id": WORKSPACE_ID}


def _put_s3_object(key: str, body: bytes) -> None:
    client = boto3.client(
        "s3",
        endpoint_url=S3_ENDPOINT,
        aws_access_key_id=os.environ.get("AWS_ACCESS_KEY_ID", "S3RVER"),
        aws_secret_access_key=os.environ.get("AWS_SECRET_ACCESS_KEY", "S3RVER"),
        region_name=os.environ.get("AWS_REGION", "us-east-1"),
        config=Config(s3={"addressing_style": "path"}),
    )
    client.put_object(Bucket=S3_BUCKET, Key=key, Body=body, ContentType=DOCX_MIME)


def _ingest(client: httpx.Client, filename: str, content: bytes, source_url: str) -> str:
    """Stage the file in object storage and start ingestion WITH a source_url.

    The public upload route has no ``source_url`` parameter (a connector supplies
    it), so this drives the ingestion service's own trigger, the same call a
    connector makes: the file lands in storage under the workspace prefix, then
    ``POST /ingest`` carries the link back to the original.
    """
    document_id = str(uuid.uuid4())
    key = f"{WORKSPACE_ID}/pilot-e2e/{document_id}/{filename}"
    _put_s3_object(key, content)
    resp = client.post(
        f"{INGESTION_URL}/ingest",
        headers={"X-API-Key": INGESTION_KEY},
        json={
            "document_id": document_id,
            "workspace_id": WORKSPACE_ID,
            "user_id": OWNER_USER,
            "filename": filename,
            "original_filename": filename,
            "content_type": DOCX_MIME,
            "size_bytes": len(content),
            "storage_backend": "s3",
            "storage_path": key,
            "storage_bucket": S3_BUCKET,
            "storage_url": f"s3://{S3_BUCKET}/{key}",
            "source_url": source_url,
        },
    )
    assert resp.status_code in (200, 202), f"ingest failed: {resp.status_code} {resp.text}"
    return document_id


def _wait_processed(client: httpx.Client, document_id: str) -> None:
    deadline = time.monotonic() + TIMEOUT
    status = None
    while time.monotonic() < deadline:
        resp = client.get(f"{API_URL}/v1/documents/{document_id}", headers=_headers(OWNER_KEY))
        if resp.status_code == 200:
            status = resp.json().get("status")
            if status == "processed":
                return
            assert status != "failed", f"document {document_id} failed ingestion: {resp.text}"
        time.sleep(3)
    pytest.fail(f"document {document_id} not processed within {TIMEOUT}s (last status={status})")


def _search(client: httpx.Client, api_key: str, query: str, limit: int = 20) -> httpx.Response:
    return client.post(
        f"{API_URL}/v1/search",
        headers={**_headers(api_key), "Content-Type": "application/json"},
        json={"query": query, "limit": limit},
    )


def _search_results(client: httpx.Client, api_key: str, query: str) -> list[dict[str, Any]]:
    resp = _search(client, api_key, query)
    assert resp.status_code == 200, f"search failed: {resp.status_code} {resp.text}"
    return resp.json()["results"]


def _wait_searchable(client: httpx.Client, document_id: str) -> None:
    """Wait until search returns the document (status flips a moment before the index does).

    Same convention as ``test_compose_integration._wait_until_searchable``: the
    vector index makes freshly stored chunks visible slightly after the document
    reads ``processed``.
    """
    _poll(
        lambda: any(
            r["document_id"] == document_id
            for r in _search_results(client, OWNER_KEY, "sides confirm arrangement reference")
        ),
        what=f"document {document_id} to become searchable",
        interval=1.0,
    )


def _poll(fn, *, what: str, timeout: float | None = None, interval: float = 2.0):
    """Call ``fn`` until it returns a truthy value; fail (not skip) on timeout."""
    deadline = time.monotonic() + (timeout or TIMEOUT)
    while time.monotonic() < deadline:
        value = fn()
        if value:
            return value
        time.sleep(interval)
    pytest.fail(f"timed out after {timeout or TIMEOUT}s waiting for {what}")


# ---------------------------------------------------------------------------
# Shared state + fixtures
# ---------------------------------------------------------------------------


@dataclass
class Flow:
    """State carried from one ordered step to the next."""

    plan: PackPlan
    mongo: Any
    doc1: str = ""
    doc2: str = ""
    # chunk_id -> search result dict for document 1 (filled by step b)
    chunks1: dict[str, dict[str, Any]] = field(default_factory=dict)
    mcp_chunk_ids: list[str] = field(default_factory=list)


def _require_stack(client: httpx.Client) -> None:
    for url in (f"{API_URL}/health", f"{INGESTION_URL}/health"):
        try:
            resp = client.get(url, timeout=5)
        except httpx.HTTPError as exc:
            pytest.skip(f"stack not reachable at {url}: {exc}")
        if resp.status_code != 200:
            pytest.skip(f"stack unhealthy at {url}: HTTP {resp.status_code}")


@pytest.fixture(scope="module")
def client() -> Iterator[httpx.Client]:
    with httpx.Client(timeout=60) as c:
        _require_stack(c)
        yield c


def _purge(client: httpx.Client) -> dict[str, Any]:
    """Start a purge of the pilot workspace and wait for its receipt."""
    start = client.post(
        f"{INGESTION_URL}/admin/workspaces/{WORKSPACE_ID}/purge",
        headers={"X-API-Key": INGESTION_KEY},
        json={"operator": "pilot-e2e"},
    )
    assert start.status_code == 202, f"purge start failed: {start.status_code} {start.text}"
    purge_id = start.json()["purge_workflow_id"]

    def _done() -> dict[str, Any] | None:
        resp = client.get(
            f"{INGESTION_URL}/admin/workspaces/{WORKSPACE_ID}/purge/{purge_id}",
            headers={"X-API-Key": INGESTION_KEY},
        )
        assert resp.status_code == 200, f"purge poll failed: {resp.status_code} {resp.text}"
        body = resp.json()
        assert body["status"] != "not_found", f"purge job vanished: {body}"
        return body if body["status"] == "completed" else None

    return _poll(_done, what="purge to complete")


@pytest.fixture(scope="module")
def flow(client: httpx.Client) -> Iterator[Flow]:
    """Provision the pilot workspace + owner key, yield the flow, clean up.

    Skips cleanly when no pack directory is configured (CI without one).
    """
    if not PACK_DIR_ENV:
        pytest.skip("E2E_VERTICAL_PACK_DIR is not set; pilot-flow E2E needs a pack directory")
    pack_root = Path(PACK_DIR_ENV)
    assert (pack_root / "vertical.yaml").is_file(), f"no vertical.yaml under {pack_root}"
    plan = _plan_from_pack(pack_root)

    mongo = MongoClient(MONGODB_URI, serverSelectionTimeoutMS=5000)
    db = mongo[MONGODB_DB]

    # Start from an empty workspace: leftovers from an aborted earlier run
    # would otherwise leak into chunk counts and reuse counters. Purge FIRST,
    # then lift the purge marker and provision: a purge revokes the workspace's
    # API keys and blocks new ingest until the marker is removed.
    _purge(client)
    _pg_clear_purge_tombstone()

    # Control plane: the workspace (owner) + the owner's workspace-scoped key.
    db.workspaces.replace_one(
        {"_id": WORKSPACE_ID},
        {"_id": WORKSPACE_ID, "user_id": OWNER_USER, "name": "Pilot E2E Workspace"},
        upsert=True,
    )
    _pg_upsert_key(OWNER_KEY_ID, OWNER_KEY, OWNER_USER)

    try:
        yield Flow(plan=plan, mongo=mongo)
    finally:
        try:
            _purge(client)
        except Exception:  # noqa: BLE001 -- best-effort cleanup, never mask the test result
            pass
        _pg_delete_keys()
        db.workspaces.delete_one({"_id": WORKSPACE_ID})
        mongo.close()


# ---------------------------------------------------------------------------
# The ordered flow. Each step depends on the previous one's state, so they run
# in file order and later steps assert their prerequisites instead of skipping.
# ---------------------------------------------------------------------------


def test_a_ingest_docx_with_source_url(client: httpx.Client, flow: Flow) -> None:
    """a. Upload a synthetic .docx with a source_url and wait for processing."""
    flow.doc1 = _ingest(
        client, f"pilot-first-{RUN}.docx", _build_docx_first(flow.plan), SOURCE_URL_1
    )
    _wait_processed(client, flow.doc1)
    _wait_searchable(client, flow.doc1)


def test_b_chunks_split_at_numbered_sections_with_tags(client: httpx.Client, flow: Flow) -> None:
    """b. One chunk per numbered section, parent heading kept, pack tags applied."""
    assert flow.doc1, "step a did not run"
    plan = flow.plan

    listed = client.get(f"{API_URL}/v1/chunks/{flow.doc1}", headers=_headers(OWNER_KEY))
    assert listed.status_code == 200, listed.text
    contents = [c["content"] for c in sorted(listed.json(), key=lambda c: c["chunk_index"])]
    # Sections 1, 1.1, 1.2, 2, 2.1 -> five chunks. "(a)"/"(b)" stay inside 1.1
    # (split_at_level=2), and the table stays inside 2.1.
    assert len(contents) == 5, f"expected 5 section chunks, got {len(contents)}: {contents}"

    def _chunk_with(text: str) -> str:
        matches = [c for c in contents if text in c]
        assert len(matches) == 1, f"{text!r} should be in exactly one chunk, got {len(matches)}"
        return matches[0]

    sub_1_1 = _chunk_with(f"A1-{RUN}")
    # Parent heading context precedes the section's own heading line.
    assert re.match(r"(?s)\s*#*\s*1\. General Provisions\s+1\.1 ", sub_1_1), sub_1_1[:200]
    # The (a)/(b) items did not start chunks of their own.
    assert SUBITEM_1 in sub_1_1 and SUBITEM_2 in sub_1_1
    assert not any(re.match(r"\s*\([a-z]\)", c) for c in contents)
    # The table rides along inside its section, rendered as a markdown table.
    sub_2_1 = _chunk_with(f"A2-{RUN}")
    assert re.match(r"(?s)\s*#*\s*2\. Additional Terms\s+2\.1 ", sub_2_1), sub_2_1[:200]
    assert f"TBLCELL-{RUN}" in sub_2_1 and "|" in sub_2_1

    # Tags + section heading come back on search results (the read path).
    results = [
        r
        for r in _search_results(client, OWNER_KEY, "sides confirm arrangement reference")
        if r["document_id"] == flow.doc1
    ]
    by_marker = {}
    for r in results:
        for marker in (f"A1-{RUN}", f"B1-{RUN}", f"A2-{RUN}"):
            if marker in r["content"]:
                by_marker[marker] = r
    assert set(by_marker) == {f"A1-{RUN}", f"B1-{RUN}", f"A2-{RUN}"}, sorted(by_marker)
    assert by_marker[f"A1-{RUN}"]["tags"][plan.filter_field] == plan.value_a
    assert by_marker[f"A2-{RUN}"]["tags"][plan.filter_field] == plan.value_a
    assert by_marker[f"B1-{RUN}"]["tags"][plan.filter_field] == plan.value_b
    assert re.match(r"#*\s*1\.1 ", by_marker[f"A1-{RUN}"]["metadata"]["section_heading"])
    assert by_marker[f"A1-{RUN}"]["source_url"] == SOURCE_URL_1
    flow.chunks1 = {r["chunk_id"]: r for r in results}


async def test_c_mcp_profile_tool_lists_and_filters(flow: Flow) -> None:
    """c. tools/list exposes the pack's profile tool; a filtered call returns its fields."""
    assert flow.chunks1, "step b did not run"
    plan = flow.plan

    async with mcp_http_session(OWNER_KEY) as session:
        listed = await session.list_tools()
        tools = {t.name: t for t in listed.tools}
        assert plan.tool_name in tools, f"{plan.tool_name!r} missing from {sorted(tools)}"
        props = tools[plan.tool_name].inputSchema["properties"]
        assert plan.filter_field in props, f"filter {plan.filter_field!r} not exposed: {props}"

        result = await session.call_tool(
            plan.tool_name,
            {"query": "sides confirm arrangement reference", plan.filter_field: plan.value_a},
        )
        assert not result.isError, result
        payload = _structured_payload(result)

    results = payload["results"]
    assert results, f"filtered profile call returned nothing: {payload}"
    for r in results:
        assert r["content"].strip(), r
        assert r["tags"][plan.filter_field] == plan.value_a, r
        assert r["section_heading"], r
        assert r["source_url"] == SOURCE_URL_1, r
    markers = {m for r in results for m in (f"A1-{RUN}", f"A2-{RUN}") if m in r["content"]}
    assert markers == {f"A1-{RUN}", f"A2-{RUN}"}, "filter must return exactly the value_a sections"
    assert not any(f"B1-{RUN}" in r["content"] for r in results), "value_b leaked through filter"

    # The MCP result carries no chunk ids; map it back through the REST result
    # so step d can compare the audited ids to what was actually returned.
    returned_markers = [m for r in results for m in (f"A1-{RUN}", f"A2-{RUN}") if m in r["content"]]
    flow.mcp_chunk_ids = [
        cid for cid, r in flow.chunks1.items() if any(m in r["content"] for m in returned_markers)
    ]
    assert len(flow.mcp_chunk_ids) == len(results)


def test_d_mcp_call_is_audited(flow: Flow) -> None:
    """d. The profile-tool call wrote an audit event in Mongo ``audit_logs``."""
    assert flow.mcp_chunk_ids, "step c did not run"
    plan = flow.plan
    logs = flow.mongo[MONGODB_DB]["audit_logs"]

    def _event() -> dict[str, Any] | None:
        return logs.find_one(
            {"workspace_id": WORKSPACE_ID, "tool_name": plan.tool_name, "surface": "mcp"}
        )

    # Audit events travel over the message queue, so they land asynchronously.
    event = _poll(_event, what="the MCP audit event to reach audit_logs", timeout=60)
    assert event["principal_type"] == "api_key"
    assert event["principal_id"] == OWNER_KEY_ID
    assert event["user_id"] == OWNER_USER
    assert event["tool_name"] == plan.tool_name and event["surface"] == "mcp"
    assert sorted(event["returned_chunk_ids"]) == sorted(flow.mcp_chunk_ids)
    assert event["result_count"] == len(flow.mcp_chunk_ids)
    assert event.get("outcome", "ok") == "ok"


def test_e_reused_section_bumps_older_chunk_reuse_count(client: httpx.Client, flow: Flow) -> None:
    """e. A newer document repeating one section increments the older chunk's reuse_count."""
    assert flow.chunks1, "step b did not run"
    marker = f"A1-{RUN}"
    before = [c for c in flow.chunks1.values() if marker in c["content"]]
    assert len(before) == 1 and before[0]["reuse_count"] == 0, before

    flow.doc2 = _ingest(
        client, f"pilot-second-{RUN}.docx", _build_docx_second(flow.plan), SOURCE_URL_2
    )
    _wait_processed(client, flow.doc2)
    _wait_searchable(client, flow.doc2)

    def _older_chunk_reuse() -> int:
        results = _search_results(client, OWNER_KEY, "sides confirm arrangement reference")
        older = [r for r in results if r["document_id"] == flow.doc1 and marker in r["content"]]
        assert len(older) == 1, older
        return older[0]["reuse_count"]

    # Detection runs inside the ingestion workflow; allow the mirrored value time to land.
    reuse = _poll(_older_chunk_reuse, what="the older chunk's reuse_count to increment")
    assert reuse == 1, f"expected exactly one reuse of the older chunk, got {reuse}"

    # The newer document's own copy is NOT counted as reused, and the untouched
    # sibling section of the older document stays at zero.
    results = _search_results(client, OWNER_KEY, "sides confirm arrangement reference")
    newer = [r for r in results if r["document_id"] == flow.doc2 and marker in r["content"]]
    assert len(newer) == 1 and newer[0]["reuse_count"] == 0
    sibling = [r for r in results if r["document_id"] == flow.doc1 and f"B1-{RUN}" in r["content"]]
    assert len(sibling) == 1 and sibling[0]["reuse_count"] == 0


def test_f_workspace_viewer_can_search_but_not_upload(client: httpx.Client, flow: Flow) -> None:
    """f. A viewer member sees the owner's documents, cannot write, and loses access when removed."""
    assert flow.doc1, "step a did not run"
    workspaces = flow.mongo[MONGODB_DB].workspaces
    _pg_upsert_key(VIEWER_KEY_ID, VIEWER_KEY, VIEWER_USER)

    # Not a member yet: the viewer's key is bound to this workspace, but the
    # workspace record does not list the viewer, so access is refused.
    denied = _search(client, VIEWER_KEY, "sides confirm arrangement reference")
    assert (
        denied.status_code == 403
    ), f"non-member must be denied: {denied.status_code} {denied.text}"

    workspaces.update_one(
        {"_id": WORKSPACE_ID},
        {
            "$set": {
                "members": [
                    {
                        "user_id": VIEWER_USER,
                        "role": "viewer",
                        "added_at": time.time(),
                        "added_by": OWNER_USER,
                    }
                ]
            }
        },
    )

    # Owner-as-tenant: the viewer reads the OWNER's documents.
    results = _search_results(client, VIEWER_KEY, "sides confirm arrangement reference")
    assert any(r["document_id"] == flow.doc1 for r in results), (
        "viewer should see the owner's document",
        [r["document_id"] for r in results],
    )

    # ...but a viewer is read-only.
    upload = client.post(
        f"{API_URL}/v1/documents",
        headers=_headers(VIEWER_KEY),
        files={
            "file": (f"viewer-{RUN}.txt", b"viewer upload attempt " + RUN.encode(), "text/plain")
        },
    )
    assert (
        upload.status_code == 403
    ), f"viewer upload must be denied: {upload.status_code} {upload.text}"
    assert "viewer" in upload.text.lower()

    # Removing the member revokes access on the very next request (no caching).
    workspaces.update_one({"_id": WORKSPACE_ID}, {"$unset": {"members": ""}})
    revoked = _search(client, VIEWER_KEY, "sides confirm arrangement reference")
    assert (
        revoked.status_code == 403
    ), f"removed member must be denied: {revoked.status_code} {revoked.text}"


def test_g_purge_leaves_no_residue(client: httpx.Client, flow: Flow) -> None:
    """g. The admin purge completes, verifies, and its receipt shows zero residue everywhere."""
    assert flow.doc1 and flow.doc2, "steps a/e did not run"
    report = _purge(client)

    assert report["verified"] is True, report
    residue = report["residue"]
    assert residue, f"purge report has no residue breakdown: {report}"
    assert all(count == 0 for count in residue.values()), f"residue left behind: {residue}"
    assert report["receipt"], f"purge produced no receipt: {report}"

    # Independent of the receipt's own accounting. The purge revoked the
    # workspace's API keys (that is part of "purge"), so the old key is refused...
    revoked = client.get(f"{API_URL}/v1/documents/{flow.doc1}", headers=_headers(OWNER_KEY))
    assert revoked.status_code == 401, f"purged workspace's key must be revoked: {revoked.text}"

    # ...and with a freshly issued key the data is really gone from the read path.
    _pg_upsert_key(OWNER_KEY_ID, OWNER_KEY, OWNER_USER)
    remaining = client.get(f"{API_URL}/v1/documents/{flow.doc1}", headers=_headers(OWNER_KEY))
    assert remaining.status_code == 404, remaining.text
    results = _search_results(client, OWNER_KEY, "sides confirm arrangement reference")
    assert not [r for r in results if r["document_id"] in (flow.doc1, flow.doc2)]
    assert flow.mongo[MONGODB_DB]["audit_logs"].count_documents({"workspace_id": WORKSPACE_ID}) == 0
